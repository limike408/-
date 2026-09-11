# -*- coding: utf-8 -*-
"""结果写入：按 附件5 模板结构用 openpyxl 填值。策略：
- 计划购电量/调整购电量：模板已预置行(2.1-12.31 共334天)，就地填 144 列。
- 充放电量/紧急购电量：模板仅示例，按天数扩展行。
"""
import os
import numpy as np
import openpyxl
import config as C


def _save(wb, path):
    try:
        wb.save(path); return path
    except PermissionError:
        alt = os.path.join(os.path.dirname(path),
                           os.path.splitext(os.path.basename(path))[0] + '_new.xlsx')
        wb.save(alt)
        print('[警告] 目标被占用，已写副本:', alt)
        return alt


def _fill_plan_sheet(ws, data):
    """data: {datetime/date日期: arr[144]}，就地填 计划/调整购电量。"""
    lookup = {}
    for k, arr in data.items():
        key = k.date() if hasattr(k, 'date') else k
        lookup[key] = arr
    for row in range(2, ws.max_row + 1):
        dt = ws.cell(row=row, column=1).value
        if dt is None or isinstance(dt, str):
            continue
        key = dt.date() if hasattr(dt, 'date') else dt
        arr = lookup.get(key)
        if arr is None:
            continue
        for t in range(C.T_IN_DAY):
            ws.cell(row=row, column=2 + t, value=round(float(arr[t]), 4))


def _clear_data_rows(ws, keep_first_col=True):
    """清空模板中残留的数据(第2行起)。
    keep_first_col=True(默认)：保留第1列(时间段标签)，只清数据列。"""
    for row in range(2, ws.max_row + 1):
        start = 2 if keep_first_col else 1
        for col in range(start, ws.max_column + 1):
            ws.cell(row=row, column=col).value = None


def _fill_chongfang(ws, dates_out, c_all, f_all, E0_all, Eend_all):
    """充放电量：日期|时间段|充电量|放电量|时刻|储电量。每日期 6块+0:00+24:00=8行。"""
    from datetime import time as tm
    _clear_data_rows(ws, keep_first_col=False)
    r = 2
    for d in dates_out:
        for h0, lab in C.BLOCK4:
            lo = h0 * 6; hi = lo + 24
            ws.cell(row=r, column=1, value=str(d))
            ws.cell(row=r, column=2, value=lab)
            ws.cell(row=r, column=3, value=round(float(np.sum(c_all[d][lo:hi])), 4))
            ws.cell(row=r, column=4, value=round(float(np.sum(f_all[d][lo:hi])), 4))
            r += 1
        ws.cell(row=r, column=1, value=str(d)); ws.cell(row=r, column=5, value='0:00')
        ws.cell(row=r, column=6, value=round(float(E0_all[d]), 4)); r += 1
        ws.cell(row=r, column=1, value=str(d)); ws.cell(row=r, column=5, value='24:00')
        ws.cell(row=r, column=6, value=round(float(Eend_all[d]), 4)); r += 1


def _fill_jinji(ws, dates_out, Q_all):
    _clear_data_rows(ws, keep_first_col=False)
    r = 2
    for d in dates_out:
        for t in range(C.T_IN_DAY):
            q = Q_all[d][t]
            if q > 1e-6:
                ws.cell(row=r, column=1, value=str(d))
                ws.cell(row=r, column=2, value=C.INTERVAL_LABELS[t])
                ws.cell(row=r, column=3, value=round(float(q), 4))
                r += 1


def write_result1(path, P, c, f, E_start, E_end, price):
    """问题1 result1.xlsx —— 共需填 2 张表。

    表1【计划购电量】：模板已预置 144 行时间段(A列2-145)，就地填 B 列购电量。
    表2【充放电量】 ：模板 7 行(2-7块 + 0:00/24:00 储电)，就地填 充电量/放电量(B/C列
                       2-7行) 与 储电量(E列，0:00在2行、24:00在3行)。
    """
    tpl = os.path.join(C.TPL_DIR, 'result1.xlsx')
    wb = openpyxl.load_workbook(tpl)

    # 表1 计划购电量：B 列(2-145) 填购电量，A 列时间段保留
    ws = wb['计划购电量']
    for t in range(C.T_IN_DAY):
        ws.cell(row=2 + t, column=2, value=round(float(P[t]), 4))

    # 表2 充放电量：就地填 6 块充电/放电量 + 0:00/24:00 储电量
    ws = wb['充放电量']
    for h0, lab in C.BLOCK4:
        lo = h0 * 6; hi = lo + 24
        r = 2 + h0 // 4                       # 块索引 0..5 -> 行 2..7
        ws.cell(row=r, column=2, value=round(float(np.sum(c[lo:hi])), 4))
        ws.cell(row=r, column=3, value=round(float(np.sum(f[lo:hi])), 4))
    ws.cell(row=2, column=4, value='0:00'); ws.cell(row=2, column=5, value=round(float(E_start), 4))
    ws.cell(row=3, column=4, value='24:00'); ws.cell(row=3, column=5, value=round(float(E_end), 4))
    wb._template = None
    return _save(wb, path)


def write_result2(path, dates_out, P, Q, c, f, E0, Eend, tpl_name='result2.xlsx'):
    """问题2/问题4-2。dates_out: 日期列表(2.1-12.31)。各 dict: date->arr/标量。"""
    tpl = os.path.join(C.TPL_DIR, tpl_name)
    wb = openpyxl.load_workbook(tpl)
    _fill_plan_sheet(wb['计划购电量'], P)
    _fill_jinji(wb['紧急购电量'], dates_out, Q)
    _fill_chongfang(wb['充放电量'], dates_out, c, f, E0, Eend)
    wb._template = None
    return _save(wb, path)


def write_result3(path, dates_out, P, A, Q, c, f, E0, Eend, tpl_name='result3.xlsx'):
    """问题3/问题4-3。"""
    tpl = os.path.join(C.TPL_DIR, tpl_name)
    wb = openpyxl.load_workbook(tpl)
    _fill_plan_sheet(wb['计划购电量'], P)
    _fill_plan_sheet(wb['调整购电量'], A)
    _fill_jinji(wb['紧急购电量'], dates_out, Q)
    _fill_chongfang(wb['充放电量'], dates_out, c, f, E0, Eend)
    wb._template = None
    return _save(wb, path)