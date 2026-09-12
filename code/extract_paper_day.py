# -*- coding: utf-8 -*-
"""抓取指定日期4天在 Q2/Q3/Q4-2/Q4-3 官方结果文件中的详细结果，供论文表格。

直接从 report/result*.xlsx 读回（与官方提交附件逐位一致，不重新求解）。
"""
import os, sys, json
sys.path.insert(0, os.path.dirname(__file__))
import numpy as np
import pandas as pd
import openpyxl
import config as C

SLOT = [m // 10 for m in C.SPEC_MINUTES]
SLOT_LAB = ['10:00-10:10', '12:00-12:10', '14:00-14:10', '16:00-16:10', '18:00-18:10', '20:00-20:10']
BLK = [(0, '0-4'), (4, '4-8'), (8, '8-12'), (12, '12-16'), (16, '16-20'), (20, '20-24')]


def blk(arr):
    return [round(float(np.sum(arr[h * 6:(h + 4) * 6])), 1) for h, _ in BLK]


def _read_day(path, sheets):
    """读回某日各 sheet 数据。sheets: {'P': 名, 'A': 名或None, 'Q': 名, 'CF': 名}。"""
    wb = openpyxl.load_workbook(path, read_only=True)
    out = {}
    for key, sn in sheets.items():
        if sn is None:
            continue
        ws = wb[sn]
        if key == 'Q':
            q = np.zeros(C.T_IN_DAY)
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None and isinstance(row[0], str) and row[0].startswith(str(d.date())):
                    lab = str(row[1])
                    t = C.INTERVAL_LABELS.index(lab)
                    q[t] = float(row[2] or 0.0)
            out['Q'] = q
        elif key == 'CF':
            e0 = e24 = None
            c = f = None
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None and isinstance(row[0], str) and row[0].startswith(str(d.date())):
                    if row[1] in ('0:00-4:00', '4:00-8:00', '8:00-12:00',
                                  '12:00-16:00', '16:00-20:00', '20:00-24:00'):
                        pass
                    if row[4] == '0:00':
                        e0 = float(row[5])
                    if row[4] == '24:00':
                        e24 = float(row[5])
            out['E0'], out['E24'] = e0, e24
        else:
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None and not isinstance(row[0], str):
                    dt = pd.to_datetime(row[0])
                    if dt.date() == d.date():
                        out[key] = np.array([float(v or 0.0) for v in row[1:1 + C.T_IN_DAY]])
                        break
    wb.close()
    return out


out = {}
for dd in C.SPEC_DAYS:
    d = pd.Timestamp(dd)
    r2 = _read_day(os.path.join(C.RES_DIR, 'result2.xlsx'),
                   dict(P='计划购电量', Q='紧急购电量', CF='充放电量'))
    r3 = _read_day(os.path.join(C.RES_DIR, 'result3.xlsx'),
                   dict(P='计划购电量', A='调整购电量', Q='紧急购电量', CF='充放电量'))
    r42 = _read_day(os.path.join(C.RES_DIR, 'result4-2.xlsx'),
                    dict(P='计划购电量', Q='紧急购电量', CF='充放电量'))
    r43 = _read_day(os.path.join(C.RES_DIR, 'result4-3.xlsx'),
                    dict(P='计划购电量', A='调整购电量', Q='紧急购电量', CF='充放电量'))

    def qsum(q):
        return round(float(q.sum()), 1)

    out[dd] = dict(
        Q2=dict(e0=r2['E0'], e24=r2['E24'], p_slots=[round(float(r2['P'][t]), 1) for t in SLOT],
                q_sum=qsum(r2['Q']), nq=int((r2['Q'] > 1e-6).sum())),
        Q4_2=dict(e0=r42['E0'], e24=r42['E24'], p_slots=[round(float(r42['P'][t]), 1) for t in SLOT],
                  q_sum=qsum(r42['Q']), nq=int((r42['Q'] > 1e-6).sum())),
        Q3=dict(e0=r3['E0'], e24=r3['E24'],
                p_slots=[round(float(r3['P'][t]), 1) for t in SLOT],
                a_slots=[round(float(r3['A'][t]), 1) for t in SLOT],
                q_sum=qsum(r3['Q']), nq=int((r3['Q'] > 1e-6).sum())),
        Q4_3=dict(e0=r43['E0'], e24=r43['E24'],
                  p_slots=[round(float(r43['P'][t]), 1) for t in SLOT],
                  a_slots=[round(float(r43['A'][t]), 1) for t in SLOT],
                  q_sum=qsum(r43['Q']), nq=int((r43['Q'] > 1e-6).sum())),
    )

print(json.dumps(out, ensure_ascii=False, indent=1))
print('SLOT_LAB', SLOT_LAB)
