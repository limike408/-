# -*- coding: utf-8 -*-
"""从结果附件(result2/3/4-2/4-3.xlsx)读回 P/A/Q，重算全年费用汇总 → report/plot_annual.csv。

用途：
1) 供 R 脚本绘制 A12（全年费用结构堆叠柱）；
2) 验收：附件数字与求解输出逐位一致（论文摘要/表格数字以本脚本输出为准）。
费用口径：
  Q2  = Σ p·P + Σ 5p·Q
  Q3  = Σ [p·min(P,A) + 0.5p(P-A)^+ + 1.5p(A-P)^+ + 5p·Q]
  Q4  = 同上，价格取附件4 当日波动电价。
"""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
import numpy as np
import pandas as pd
import openpyxl
import config as C
import data

# 附件5 模板 sheet 名
S_PLAN, S_ADJ, S_JINJI = '计划购电量', '调整购电量', '紧急购电量'


def _read_day_matrix(path, sheet):
    """读回 计划/调整购电量 sheet → {date: arr144}。"""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[sheet]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None or isinstance(row[0], str):
            continue
        dt = pd.to_datetime(row[0])
        vals = [float(v) if v is not None else 0.0 for v in row[1:1 + C.T_IN_DAY]]
        out[dt] = np.array(vals)
    wb.close()
    return out


def _read_jinji(path):
    """读回 紧急购电量 sheet → {date: {t: q}}。"""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[S_JINJI]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None or not isinstance(row[0], str):
            continue
        try:
            dt = pd.to_datetime(row[0])
        except Exception:
            continue
        lab = str(row[1])
        q = float(row[2]) if row[2] is not None else 0.0
        t = C.INTERVAL_LABELS.index(lab)
        out.setdefault(dt, {})[t] = q
    wb.close()
    return out


def _q_vec(jinji_day, n=144):
    q = np.zeros(n)
    for t, v in jinji_day.items():
        q[t] = v
    return q


def _q2_fees(path, price_map):
    """Q2/Q4-2：Σp·P + Σ5p·Q。price_map: {date: arr144}。"""
    P = _read_day_matrix(path, S_PLAN)
    J = _read_jinji(path)
    plan = emerg = 0.0
    for d, arr in P.items():
        pr = price_map[pd.Timestamp(d)]
        plan += float(np.sum(pr * arr))
        emerg += float(np.sum(5.0 * pr * _q_vec(J.get(pd.Timestamp(d), {}))))
    return plan, emerg


def _q3_fees(path, price_map):
    """Q3/Q4-3：Σ[p·min(P,A) + 0.5p(P-A)^+ + 1.5p(A-P)^+ + 5p·Q]。"""
    P = _read_day_matrix(path, S_PLAN)
    A = _read_day_matrix(path, S_ADJ)
    J = _read_jinji(path)
    grid = emerg = 0.0
    for d, pv in P.items():
        dt = pd.Timestamp(d)
        pr = price_map[dt]
        av = A[dt]
        u = np.maximum(pv - av, 0.0)
        w = np.maximum(av - pv, 0.0)
        grid += float(np.sum(pr * np.minimum(pv, av) + 0.5 * pr * u + 1.5 * pr * w))
        emerg += float(np.sum(5.0 * pr * _q_vec(J.get(dt, {}))))
    return grid, emerg


def main():
    a1 = data.load_attach1()
    fp = a1['price']
    dates4, p4 = data.load_attach4()
    price_fix = {d: fp for d in dates4}
    price_var = {d: p4[i] for i, d in enumerate(dates4)}

    g2, e2 = _q2_fees(os.path.join(C.RES_DIR, 'result2.xlsx'), price_fix)
    g3, e3 = _q3_fees(os.path.join(C.RES_DIR, 'result3.xlsx'), price_fix)
    g42, e42 = _q2_fees(os.path.join(C.RES_DIR, 'result4-2.xlsx'), price_var)
    g43, e43 = _q3_fees(os.path.join(C.RES_DIR, 'result4-3.xlsx'), price_var)

    rows = [
        dict(label='问题二\n(预报式两阶段随机)', grid_fee=round(g2, 2), emerg_fee=round(e2, 2)),
        dict(label='问题三\n(滚动随机MPC)', grid_fee=round(g3, 2), emerg_fee=round(e3, 2)),
        dict(label='问题四·问题三\n(波动电价)', grid_fee=round(g43, 2), emerg_fee=round(e43, 2)),
    ]
    df = pd.DataFrame(rows)
    df.to_csv(os.path.join(C.RES_DIR, 'plot_annual.csv'), index=False, encoding='utf-8-sig')
    print('[plot_annual.csv 已写]')
    print(df.to_string(index=False))
    print(f'\n[附件读回汇总] Q2 合计={g2+e2:,.2f} | Q3 合计={g3+e3:,.2f} | '
          f'Q4-2 合计={g42+e42:,.2f} | Q4-3 合计={g43+e43:,.2f}')
    print(f'[对比] Q4-2 较 Q2 上浮 {(g42+e42)/(g2+e2)*100-100:+.2f}% | '
          f'Q4-3 较 Q3 上浮 {(g43+e43)/(g3+e3)*100-100:+.2f}%')
    return dict(g2=g2, e2=e2, g3=g3, e3=e3, g42=g42, e42=e42, g43=g43, e43=e43)


if __name__ == '__main__':
    main()
